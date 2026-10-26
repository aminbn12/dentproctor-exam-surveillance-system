from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy.orm import Session, selectinload
from typing import List
from app.database import get_db
from app.models.exam import Exam
from app.models.room import Room
from app.schemas.exam import ExamResponse, ExamCreate, ExamUpdate
from app.middleware.auth import get_admin_user
from app.models.user import User

router = APIRouter(prefix="/api/exams", tags=["exams"])

@router.get("", response_model=List[ExamResponse])
async def get_exams(
    db: Session = Depends(get_db),
    skip: int = 0,
    limit: int = 1000
):
    """Récupérer tous les examens (avec pagination et préchargement des salles)"""
    exams = db.query(Exam).options(
        selectinload(Exam.rooms)
    ).offset(skip).limit(limit).all()
    return exams

@router.post("", response_model=ExamResponse)
async def create_exam(
    exam: ExamCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_admin_user)
):
    """Créer un examen (avec association des salles)"""
    # 1. Valider que toutes les salles existent (si fournies)
    rooms_to_associate: List[Room] = []
    if exam.room_ids:
        for room_id in exam.room_ids:
            room = db.query(Room).filter(Room.id == room_id).first()
            if not room:
                raise HTTPException(
                    status_code=404,
                    detail=f"Salle avec ID '{room_id}' non trouvée"
                )
            rooms_to_associate.append(room)

    # 2. Créer l'examen
    new_exam = Exam(
        date=exam.date,
        time=exam.time,
        duration=exam.duration,
        promo=exam.promo,
        subject=exam.subject
    )
    db.add(new_exam)
    db.commit()
    db.refresh(new_exam)

    # 3. Associer les salles validées
    for room in rooms_to_associate:
        new_exam.rooms.append(room)
    db.commit()
    db.refresh(new_exam)

    return new_exam

@router.put("/{exam_id}", response_model=ExamResponse)
async def update_exam(
    exam_id: str,
    exam_data: ExamUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_admin_user)
):
    """Mettre à jour un examen"""
    exam = db.query(Exam).filter(Exam.id == exam_id).first()
    if not exam:
        raise HTTPException(status_code=404, detail="Examen non trouvé")
    
    data_dict = exam_data.dict(exclude_unset=True)
    room_ids = data_dict.pop("room_ids", None)
    
    for key, value in data_dict.items():
        setattr(exam, key, value)
    
    if room_ids is not None:
        rooms = db.query(Room).filter(Room.id.in_(room_ids)).all()
        exam.rooms = rooms
    
    db.commit()
    db.refresh(exam)
    return exam

@router.delete("/{exam_id}")
async def delete_exam(
    exam_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_admin_user)
):
    """Supprimer un examen"""
    exam = db.query(Exam).filter(Exam.id == exam_id).first()
    if not exam:
        raise HTTPException(status_code=404, detail="Examen non trouvé")
    
    db.delete(exam)
    db.commit()
    return {"message": "✅ Examen supprimé"}


from fastapi.responses import Response
from pydantic import BaseModel
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
from datetime import datetime
import os
import re

class MacroRowItem(BaseModel):
    date: str = ""
    time: str = ""
    endTime: str = ""
    promo: str = ""
    subject: str = ""
    room: str = ""
    profs: str = ""
    residents: str = ""
    adm: str = ""
    it: str = ""
    zone: str = ""

class MacroExportRequest(BaseModel):
    rows: List[MacroRowItem]

FRENCH_DAYS = ["lundi", "mardi", "mercredi", "jeudi", "vendredi", "samedi", "dimanche"]

def format_french_date(iso_date: str) -> str:
    try:
        dt = datetime.strptime(iso_date.strip(), "%Y-%m-%d")
        day_name = FRENCH_DAYS[dt.weekday()]
        return f"{day_name} {dt.strftime('%d/%m/%Y')}"
    except Exception:
        return iso_date.strip()

PROMO_COLORS = {
    "FM6MD1": "FEF9C3", # yellow-100
    "FM6MD2": "DBEAFE", # blue-100
    "FM6MD3": "DCFCE7", # green-100
    "FM6MD4": "FFE4E6", # rose-100
    "FM6MD5": "F3E8FF", # purple-100
}

@router.post("/export-macro-xlsm")
async def export_macro_xlsm(
    payload: MacroExportRequest
):
    """Générer le classeur .xlsm officiel pré-rempli, stylisé aux couleurs de l'export normal avec colonnes de dates garanties visibles"""
    base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
    template_path = os.path.join(base_dir, "public", "Modele_Planning_Surveillance.xlsm")
    if not os.path.exists(template_path):
        template_path = os.path.join(base_dir, "dist", "Modele_Planning_Surveillance.xlsm")
    if not os.path.exists(template_path):
        template_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "..", "public", "Modele_Planning_Surveillance.xlsm"))
    
    if not os.path.exists(template_path):
        raise HTTPException(status_code=404, detail="Modèle .xlsm introuvable")
    
    wb = openpyxl.load_workbook(template_path, keep_vba=True)
    ws = wb["Feuil1"]
    
    # 1. Conserver la ligne 1 (en-têtes) et vider les lignes de données précédentes
    if ws.max_row > 1:
        ws.delete_rows(2, amount=ws.max_row)
        
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB'),
    )

    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Style de l'en-tête Feuil1
    ws.row_dimensions[1].height = 28
    for col_idx in range(1, 12):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border

    # 2. Injecter les données avec styles identiques à l'export Excel classique (alternance jour blanc / gris très clair #F1F5F9)
    unique_dates_raw = []
    for r in payload.rows:
        if r.date and r.date not in unique_dates_raw:
            unique_dates_raw.append(r.date)

    gray_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    for r_idx, r in enumerate(payload.rows, start=2):
        row_values = [
            r.date, r.time, r.endTime, r.promo, r.subject, r.room,
            r.profs, r.residents, r.adm, r.it, r.zone
        ]
        ws.append(row_values)
        ws.row_dimensions[r_idx].height = 22
        
        date_idx = unique_dates_raw.index(r.date) if r.date in unique_dates_raw else 0
        is_alt_day = (date_idx % 2 == 1)
        day_fill = gray_fill if is_alt_day else white_fill

        bg_promo = PROMO_COLORS.get(r.promo, "FFFFFF")
        promo_fill = PatternFill(start_color=bg_promo, end_color=bg_promo, fill_type="solid")
        
        for col_idx in range(1, 12):
            cell = ws.cell(row=r_idx, column=col_idx)
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10.5)
            
            if col_idx in [1, 2, 3]:  # Date, Heures
                cell.fill = day_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 4:  # Promo (garde sa couleur de promo)
                cell.fill = promo_fill
                cell.font = Font(name="Calibri", size=10.5, bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 5:  # Épreuve (garde sa couleur de promo)
                cell.fill = promo_fill
                cell.font = Font(name="Calibri", size=10.5, bold=True)
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx == 6:  # Salle (Rouge vif et fond du jour)
                cell.fill = day_fill
                cell.font = Font(name="Calibri", size=10.5, bold=True, color="DC2626")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:  # Surveillants (fond du jour)
                cell.fill = day_fill
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Largeurs généreuses pour Feuil1
    col_widths = {
        "A": 15, "B": 13, "C": 13, "D": 14, "E": 32, "F": 18,
        "G": 32, "H": 32, "I": 25, "J": 20, "K": 22
    }
    for col, w in col_widths.items():
        ws.column_dimensions[col].width = w

    ws.views.sheetView[0].showGridLines = True

    # 3. Pré-génération et mise en forme des 5 feuilles d'analyse
    SHEETS_CONFIG = [
        ("Enseignants", lambda row: row.profs),
        ("Residents", lambda row: row.residents),
        ("Administration", lambda row: row.adm),
        ("IT", lambda row: row.it),
        ("RespZones", lambda row: row.zone)
    ]

    unique_dates = []
    for r in payload.rows:
        d = format_french_date(r.date)
        if d not in unique_dates:
            unique_dates.append(d)

    indigo_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    indigo_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    sub_fill = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")
    sub_font = Font(name="Calibri", size=10, bold=True, color="1E1B4B")
    amber_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    amber_font = Font(name="Calibri", size=11, bold=True, color="78350F")

    for sheet_name, get_staff_fn in SHEETS_CONFIG:
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        calc_ws = wb.create_sheet(title=sheet_name)
        
        calc_ws.views.sheetView[0].showGridLines = True
        
        # Ligne 1 : Bandeau de titre
        calc_ws.row_dimensions[1].height = 36
        title_cell = calc_ws.cell(row=1, column=1, value=f"PLANNING - {sheet_name.upper()}")
        title_cell.font = Font(name="Calibri", size=13, bold=True, color="1E1B4B")
        title_cell.alignment = Alignment(horizontal="left", vertical="center")
        
        # Ligne 2 & 3 : En-têtes de tableau
        calc_ws.row_dimensions[2].height = 26
        calc_ws.row_dimensions[3].height = 20
        
        calc_ws.cell(row=2, column=1, value="CATEGORIE")
        calc_ws.cell(row=3, column=1, value=f"Surveillant ({sheet_name})")
        
        c = 2
        for d in unique_dates:
            calc_ws.cell(row=2, column=c, value=d)
            calc_ws.merge_cells(start_row=2, start_column=c, end_row=2, end_column=c+1)
            calc_ws.cell(row=3, column=c, value="Matin")
            calc_ws.cell(row=3, column=c+1, value="Aprem")
            c += 2
            
        calc_ws.cell(row=2, column=c, value="TOTAL")
        calc_ws.cell(row=3, column=c, value="General")
        calc_ws.merge_cells(start_row=2, start_column=c, end_row=3, end_column=c)
        total_cols = c
        
        # Agréger le personnel
        person_counts = {}
        # Pour les enseignants : stocker si la personne est responsable de promo
        person_resp_info = {}  # clean_nom -> {"digit": str, "color": str}

        for r in payload.rows:
            d = format_french_date(r.date)
            try:
                h_val = int(r.time.split(":")[0])
            except:
                h_val = 9
            period = "Matin" if h_val < 13 else "Aprem"
            names_str = get_staff_fn(r) or ""
            for n in names_str.replace("\n", ",").split(","):
                raw_nom = n.strip()
                resp_match = re.search(r'\(Resp(?:\s*(\d+))?\)', raw_nom, re.IGNORECASE)
                is_resp = bool(resp_match)
                clean_nom = re.sub(r'\(Resp(?:\s*\d*)?\)', '', raw_nom, flags=re.IGNORECASE).strip()
                if not clean_nom:
                    continue
                if clean_nom not in person_counts:
                    person_counts[clean_nom] = {}
                k_slot = (d, period)
                person_counts[clean_nom][k_slot] = person_counts[clean_nom].get(k_slot, 0) + 1

                if sheet_name == "Enseignants" and is_resp and clean_nom not in person_resp_info:
                    digit = resp_match.group(1) if (resp_match and resp_match.group(1)) else None
                    if not digit:
                        m_p = re.search(r'\d+$', r.promo or '')
                        digit = m_p.group(0) if m_p else (r.promo[-1:] if r.promo else "")
                    if not digit:
                        digit = "1"
                    
                    p_color = PROMO_COLORS.get(r.promo, "")
                    if not p_color:
                        p_color = PROMO_COLORS.get(f"FM6MD{digit}", "FEF9C3")
                    
                    person_resp_info[clean_nom] = {
                        "digit": digit,
                        "color": p_color
                    }

        r_idx = 4
        for p_name in sorted(person_counts.keys()):
            resp_info = person_resp_info.get(p_name) if sheet_name == "Enseignants" else None
            if resp_info:
                display_name = f"{p_name} (Resp {resp_info['digit']})"
                c1 = calc_ws.cell(row=r_idx, column=1, value=display_name)
                c_color = resp_info["color"]
                c1.fill = PatternFill(start_color=c_color, end_color=c_color, fill_type="solid")
                c1.font = Font(name="Calibri", size=11, bold=True, color="1E1B4B")
            else:
                c1 = calc_ws.cell(row=r_idx, column=1, value=p_name)
                c1.font = Font(name="Calibri", size=11, bold=True, color="1E293B")
            
            total_p = 0
            c = 2
            for d in unique_dates:
                v_m = person_counts[p_name].get((d, "Matin"), 0)
                v_a = person_counts[p_name].get((d, "Aprem"), 0)
                
                # Matin standard
                cell_m = calc_ws.cell(row=r_idx, column=c, value="-" if v_m == 0 else v_m)
                if v_m > 0:
                    cell_m.font = Font(name="Calibri", size=11, bold=True)
                
                # Aprem standard
                cell_a = calc_ws.cell(row=r_idx, column=c+1, value="-" if v_a == 0 else v_a)
                if v_a > 0:
                    cell_a.font = Font(name="Calibri", size=11, bold=True)
                
                total_p += v_m + v_a
                c += 2
            total_cell = calc_ws.cell(row=r_idx, column=c, value=total_p)
            total_cell.font = Font(name="Calibri", size=11, bold=True)
            r_idx += 1
            
        if r_idx == 4:
            calc_ws.cell(row=4, column=1, value=f"(Aucune affectation détectée pour {sheet_name})")
            r_idx = 5
            
        # Mise en forme Lignes 2 et 3
        for c_i in range(1, total_cols + 1):
            cell_2 = calc_ws.cell(row=2, column=c_i)
            cell_2.fill = indigo_fill
            cell_2.font = indigo_font
            cell_2.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
            cell_2.border = thin_border
            
            cell_3 = calc_ws.cell(row=3, column=c_i)
            if c_i != total_cols:
                cell_3.fill = sub_fill
                cell_3.font = sub_font
            else:
                cell_3.fill = indigo_fill
                cell_3.font = indigo_font
            cell_3.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
            cell_3.border = thin_border

        for row_i in range(4, r_idx):
            calc_ws.row_dimensions[row_i].height = 20
            for c_i in range(1, total_cols + 1):
                cell = calc_ws.cell(row=row_i, column=c_i)
                cell.border = thin_border
                if c_i == 1:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif c_i == total_cols:
                    cell.fill = amber_fill
                    cell.font = amber_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

        # Largeurs de colonnes : GARANTIE D'AFFICHAGE DU JOUR ET DE LA DATE SANS TRONCATURE
        calc_ws.column_dimensions["A"].width = 32
        for col_i in range(2, total_cols):
            col_letter = get_column_letter(col_i)
            calc_ws.column_dimensions[col_letter].width = 13.5
        calc_ws.column_dimensions[get_column_letter(total_cols)].width = 13

    buf = BytesIO()
    wb.save(buf)
    content = buf.getvalue()
    try:
        wb.close()
    except Exception:
        pass
    
    return Response(
        content=content,
        media_type="application/vnd.ms-excel.sheet.macroEnabled.12",
        headers={
            "Content-Disposition": "attachment; filename=Planning_Surveillances_Macro.xlsm"
        }
    )


